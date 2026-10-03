"""
demo_routes.py
--------------
Rutas para el modo Demo: simula ventas físicas desde un Excel y las descuenta
en inventario de Shopify en tiempo real.

Productos demo predefinidos (SKUs asignados al configurar la app):
  DEMO-001  →  producto 1-Demo
  DEMO-002  →  Producto -2 Demo
  DEMO-003  →  Producto -3 Demo
"""
import io
import logging
from typing import Any, Dict, List

import httpx
import openpyxl
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from fastapi.responses import StreamingResponse

from config import settings

logger = logging.getLogger("inventory_sync.demo")

router = APIRouter(prefix="/demo", tags=["Demo"])

# ---------------------------------------------------------------------------
# Productos demo hardcodeados (se leen también desde Shopify en vivo)
# ---------------------------------------------------------------------------
DEMO_PRODUCTS = [
    {"sku": "DEMO-001", "nombre": "Producto 1 - Demo", "inventory_item_id": "45849171263555"},
    {"sku": "DEMO-002", "nombre": "Producto 2 - Demo", "inventory_item_id": "45851129086019"},
    {"sku": "DEMO-003", "nombre": "Producto 3 - Demo", "inventory_item_id": "45851129118787"},
]

SHOPIFY_LOCATION_GID = "gid://shopify/Location/80449372227"


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _shopify_headers() -> Dict[str, str]:
    token = settings.SHOPIFY_ACCESS_TOKEN or ""
    return {
        "X-Shopify-Access-Token": token.strip(),
        "Content-Type": "application/json",
    }


def _shopify_url() -> str:
    domain = (settings.SHOP_DOMAIN or "").replace("https://", "").replace("http://", "").strip("/")
    version = getattr(settings, "SHOPIFY_API_VERSION", "2024-01")
    return f"https://{domain}/admin/api/{version}/graphql.json"


async def _get_shopify_stock() -> List[Dict[str, Any]]:
    """Consulta el inventario actual de los 3 productos demo en Shopify."""
    inv_ids = [f"gid://shopify/InventoryItem/{p['inventory_item_id']}" for p in DEMO_PRODUCTS]
    query = """
    query getInventory($ids: [ID!]!) {
        nodes(ids: $ids) {
            ... on InventoryItem {
                id
                sku
                inventoryLevel(locationId: "%s") {
                    quantities(names: ["available"]) {
                        name
                        quantity
                    }
                }
            }
        }
    }
    """ % SHOPIFY_LOCATION_GID

    async with httpx.AsyncClient(timeout=15.0) as client:
        r = await client.post(
            _shopify_url(),
            headers=_shopify_headers(),
            json={"query": query, "variables": {"ids": inv_ids}},
        )
        r.raise_for_status()
        nodes = r.json().get("data", {}).get("nodes", [])

    result = []
    for i, node in enumerate(nodes):
        if not node:
            continue
        product = DEMO_PRODUCTS[i]
        level = node.get("inventoryLevel") or {}
        quantities = level.get("quantities", [])
        qty = quantities[0].get("quantity", 0) if quantities else 0
        result.append({
            "sku": product["sku"],
            "nombre": product["nombre"],
            "inventory_item_id": product["inventory_item_id"],
            "stock_actual": qty,
        })
    return result


async def _adjust_shopify_inventory(inventory_item_id: str, delta: int) -> Dict[str, Any]:
    """
    Ajusta el inventario de un item en Shopify con una cantidad relativa (negativa = descuento).
    Usa inventoryAdjustQuantities mutation (API 2024-01+).
    """
    mutation = """
    mutation adjustInventory($input: InventoryAdjustQuantitiesInput!) {
        inventoryAdjustQuantities(input: $input) {
            inventoryAdjustmentGroup {
                createdAt
                reason
                changes {
                    name
                    delta
                    quantityAfterChange
                }
            }
            userErrors {
                field
                message
            }
        }
    }
    """
    variables = {
        "input": {
            "reason": "correction",
            "name": "available",
            "changes": [
                {
                    "inventoryItemId": f"gid://shopify/InventoryItem/{inventory_item_id}",
                    "locationId": SHOPIFY_LOCATION_GID,
                    "delta": delta,
                }
            ],
        }
    }

    async with httpx.AsyncClient(timeout=15.0) as client:
        r = await client.post(
            _shopify_url(),
            headers=_shopify_headers(),
            json={"query": mutation, "variables": variables},
        )
        r.raise_for_status()
        data = r.json()

    errors = (
        data.get("data", {})
        .get("inventoryAdjustQuantities", {})
        .get("userErrors", [])
    )
    if errors:
        raise HTTPException(status_code=400, detail=f"Shopify error: {errors[0]['message']}")

    changes = (
        data.get("data", {})
        .get("inventoryAdjustQuantities", {})
        .get("inventoryAdjustmentGroup", {})
        .get("changes", [])
    )
    qty_after = changes[0].get("quantityAfterChange", "?") if changes else "?"
    return {"inventory_item_id": inventory_item_id, "delta": delta, "stock_resultante": qty_after}


# ---------------------------------------------------------------------------
# Endpoints
# ---------------------------------------------------------------------------

@router.get("/products")
async def get_demo_products():
    """
    Devuelve los 3 productos demo con su stock actual en Shopify (tiempo real).
    """
    try:
        products = await _get_shopify_stock()
        return {"success": True, "products": products}
    except Exception as e:
        logger.error(f"[DEMO] Error consultando stock Shopify: {e}")
        return {"success": False, "message": str(e), "products": DEMO_PRODUCTS}


@router.get("/template")
def download_excel_template():
    """
    Descarga una plantilla Excel con los 3 productos demo lista para llenar.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ventas Demo"

    # Estilos de encabezado
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    header_fill = PatternFill("solid", fgColor="1A1A2E")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    headers = ["SKU", "Producto", "Stock Actual (Shopify)", "Unidades Vendidas", "Stock Resultante"]
    col_widths = [14, 30, 24, 22, 20]

    for col, (h, w) in enumerate(zip(headers, col_widths), 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
        ws.column_dimensions[get_column_letter(col)].width = w

    ws.row_dimensions[1].height = 22

    # Filas de productos demo (con fórmula para stock resultante)
    product_fill = PatternFill("solid", fgColor="F0F4FF")
    for row, p in enumerate(DEMO_PRODUCTS, 2):
        ws.cell(row=row, column=1, value=p["sku"]).border = border
        ws.cell(row=row, column=2, value=p["nombre"]).border = border
        stock_cell = ws.cell(row=row, column=3, value="(consultar)")
        stock_cell.border = border
        stock_cell.alignment = Alignment(horizontal="center")

        venta_cell = ws.cell(row=row, column=4, value=0)
        venta_cell.border = border
        venta_cell.alignment = Alignment(horizontal="center")
        venta_cell.fill = PatternFill("solid", fgColor="FFF9C4")  # amarillo = editable

        result_cell = ws.cell(row=row, column=5, value=f"=C{row}-D{row}")
        result_cell.border = border
        result_cell.alignment = Alignment(horizontal="center")
        result_cell.fill = product_fill

    # Instrucciones
    ws.cell(row=6, column=1, value="ℹ️ Instrucciones:").font = Font(bold=True)
    ws.cell(row=7, column=1, value="1. Llena la columna 'Unidades Vendidas' (columna D) con las ventas del día.")
    ws.cell(row=8, column=1, value="2. Sube este archivo en la plataforma (sección Demo).")
    ws.cell(row=9, column=1, value="3. El sistema descontará automáticamente el inventario en Shopify.")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=demo_ventas.xlsx"},
    )


@router.post("/sync-excel")
async def sync_excel_to_shopify(file: UploadFile = File(...)):
    """
    Recibe un Excel con ventas (columna D = unidades vendidas por SKU),
    descuenta el inventario en Shopify en tiempo real y devuelve el resultado.
    """
    if not file.filename or not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="Sube un archivo Excel (.xlsx o .xls).")

    content = await file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        ws = wb.active
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo leer el Excel: {e}")

    # Mapa SKU → inventory_item_id
    sku_map = {p["sku"]: p["inventory_item_id"] for p in DEMO_PRODUCTS}

    sales: List[Dict[str, Any]] = []
    errors: List[str] = []

    # Leer filas desde fila 2 (fila 1 = encabezados)
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        sku = str(row[0]).strip().upper()
        try:
            units_sold = int(row[3] or 0)
        except (ValueError, TypeError):
            units_sold = 0

        if sku not in sku_map:
            continue  # ignora SKUs que no son demo
        if units_sold <= 0:
            continue  # sin venta, nada que hacer

        sales.append({"sku": sku, "inventory_item_id": sku_map[sku], "units_sold": units_sold})

    if not sales:
        return {
            "success": False,
            "message": "No se encontraron ventas en el Excel. Llena la columna 'Unidades Vendidas' (columna D).",
            "results": [],
        }

    # Aplicar descuentos en Shopify
    results = []
    for sale in sales:
        try:
            res = await _adjust_shopify_inventory(
                inventory_item_id=sale["inventory_item_id"],
                delta=-sale["units_sold"],  # negativo = descontar
            )
            results.append({
                "sku": sale["sku"],
                "units_sold": sale["units_sold"],
                "stock_resultante": res["stock_resultante"],
                "status": "✅ Descontado en Shopify",
            })
            logger.info(f"[DEMO] {sale['sku']}: -{sale['units_sold']} → stock={res['stock_resultante']}")
        except Exception as e:
            errors.append(f"{sale['sku']}: {e}")
            results.append({
                "sku": sale["sku"],
                "units_sold": sale["units_sold"],
                "stock_resultante": "?",
                "status": f"❌ Error: {e}",
            })

    return {
        "success": len(errors) == 0,
        "message": f"Se procesaron {len(results)} productos. {len(errors)} errores." if errors
                   else f"✅ {len(results)} productos descontados en Shopify exitosamente.",
        "results": results,
        "errors": errors,
    }


@router.post("/manual-sale")
async def manual_sale(body: Dict[str, Any]):
    """
    Registra una venta manual (sin Excel) para un SKU demo.
    Body: { "sku": "DEMO-001", "units_sold": 5 }
    """
    sku = str(body.get("sku", "")).strip().upper()
    units_sold = int(body.get("units_sold", 0))

    sku_map = {p["sku"]: p["inventory_item_id"] for p in DEMO_PRODUCTS}
    if sku not in sku_map:
        raise HTTPException(status_code=400, detail=f"SKU '{sku}' no encontrado en productos demo.")
    if units_sold <= 0:
        raise HTTPException(status_code=400, detail="Ingresa un número de unidades mayor a 0.")

    try:
        res = await _adjust_shopify_inventory(sku_map[sku], delta=-units_sold)
        logger.info(f"[DEMO MANUAL] {sku}: -{units_sold} → stock={res['stock_resultante']}")
        return {
            "success": True,
            "sku": sku,
            "units_sold": units_sold,
            "stock_resultante": res["stock_resultante"],
            "message": f"✅ Venta registrada. Stock de {sku} actualizado a {res['stock_resultante']} en Shopify.",
        }
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
